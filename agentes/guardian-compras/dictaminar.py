"""Dictamen semanal de órdenes de compra · Kabla Comercial.

Uso: python3 dictaminar.py lineas.json AAAA-MM-DD [carpeta_salida] [logo.png]
lineas.json: lista de filas devueltas por sql/lineas_oc.sql.
Reglas: ver .claude/skills/guardian-compras/SKILL.md.
"""
import json, math, sys, os
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image

FACTOR_COSTO_PUESTO = 1.12
COBERTURA_MAX_DIAS = 180
GASTOS = 0.15
MARGEN_MIN_ANALIZADOR = 0.10
BASICO = {"De Linea - Basico", "De Linea", "De Linea - Variable"}
BAJO_PEDIDO = {"Bajo pedido", "Bajo pedido especial", "Bajo Proyecto"}
SIN_COMPRA = {"Descontinuado", "Hasta Agotar Existencias", "Outlet"}
LIMITE_FIRMA_USD = 5000  # OC aprobadas completas debajo de este monto las firma el gerente de compras
TIPO_CAMBIO = float(os.environ.get("TIPO_CAMBIO_USD", "18.5"))

ROYAL, ONYX, SEASHELL, PLATINIUM, GINGER_T, ROYAL_T, INK = "FF8000", "444444", "FFF6F2", "DEDEDE", "FFDECB", "FBD7BD", "AD5600"


def n(x):
    try:
        return float(x) if x not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def parse_fecha(txt):
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(txt), fmt).date()
        except (TypeError, ValueError):
            pass
    return None


HOY = date.today()


def dictaminar(l):
    cant, precio = n(l.get("cant")), n(l.get("precio"))
    stock, transito = n(l.get("stock_vendible")), n(l.get("en_transito"))
    venta, pedidos, dist_iv = n(l.get("venta_12m")), n(l.get("pedidos_abiertos")), n(l.get("dist_iv"))
    tipo = l.get("tipo_producto") or "Sin tipo"
    articulo = l.get("articulo") or ""
    vd = venta / 365 if venta > 0 else 0
    aprob, decision, razones, notas = cant, "APROBAR", [], []

    if tipo in SIN_COMPRA:
        aprob, decision = 0, "RECHAZAR"
        razones.append(f"Artículo {tipo.lower()}: no se compra.")
    elif tipo in BAJO_PEDIDO:
        cap = max(0, math.floor(pedidos - stock - transito))
        if cap <= 0:
            aprob, decision = 0, "RECHAZAR"
            razones.append(f"{tipo}: no hay pedidos de cliente abiertos (6 meses) que justifiquen la compra; stock {stock:,.0f}, en tránsito {transito:,.0f}. Si es para licitación o contrato, Compras debe ligar la OC al pedido.")
        elif cap < cant:
            aprob, decision = cap, "REDUCIR"
            razones.append(f"{tipo}: solo {cap:,.0f} piezas están respaldadas por pedidos de cliente abiertos.")
    elif tipo in BASICO:
        if vd == 0:
            aprob, decision = 0, "RECHAZAR"
            razones.append("Producto básico sin venta en 12 meses.")
        else:
            llega = parse_fecha(l.get("llegada"))
            dias_llegada = max(0, (llega - HOY).days) if llega else 0
            proyectado = max(0.0, stock + transito - vd * dias_llegada)
            cap = math.floor(COBERTURA_MAX_DIAS * vd - proyectado)
            dias = (proyectado + cant) / vd
            if cap <= 0:
                aprob, decision = 0, "RECHAZAR"
                razones.append(f"Al llegar ({dias_llegada} días) todavía habrá {proyectado / vd:,.0f} días de inventario entre stock y tránsito; el tope es {COBERTURA_MAX_DIAS}.")
            elif cap < cant:
                aprob, decision = cap, "REDUCIR"
                razones.append(f"Al llegar quedarían {dias:,.0f} días de inventario; {cap:,.0f} piezas dejan {COBERTURA_MAX_DIAS}.")
    elif tipo == "De Linea - Emergente":
        notas.append("Emergente: sin tope de cobertura.")
    else:
        notas.append(f"Tipo '{tipo}': sin regla de cobertura.")

    es_analizador = articulo.lower().startswith(("analizador", "contador"))
    if es_analizador:
        notas.append("Analizador: la venta facturada no incluye equipos colocados en comodato; validar demanda con Comercial antes de firmar.")
    if es_analizador and precio > 0 and decision != "RECHAZAR":
        if dist_iv <= 0:
            notas.append("Analizador sin precio DISTRIBUIDOR IV: margen no verificable, revisar.")
        else:
            margen = (dist_iv - precio * FACTOR_COSTO_PUESTO - GASTOS * dist_iv) / dist_iv
            if margen < MARGEN_MIN_ANALIZADOR:
                tope = dist_iv * (1 - GASTOS - MARGEN_MIN_ANALIZADOR) / FACTOR_COSTO_PUESTO
                aprob, decision = 0, "RECHAZAR"
                razones.append(f"Margen después de gastos de {margen:.1%} (mínimo 10%). Precio de compra máximo que cumple: ${tope:,.0f}.")
            else:
                notas.append(f"Margen después de gastos: {margen:.1%}.")

    mejor = n(l.get("mejor_precio_12m"))
    renegociar = 0.0
    if precio > 0 and mejor > 0 and precio > mejor * 1.005 and aprob > 0:
        renegociar = (precio - mejor) * aprob
        notas.append(f"Renegociar: pagamos ${mejor:,.2f} a este proveedor en 12 meses; diferencia ${renegociar:,.0f}.")
    if precio == 0:
        notas.append("Línea sin costo (bonificación).")
    if not razones and decision == "APROBAR":
        razones.append("Cumple las reglas.")

    return {
        "oc": l.get("oc"), "proveedor": l.get("proveedor"), "articulo": articulo, "tipo": tipo,
        "llegada": l.get("llegada"), "cant": cant, "precio": precio, "importe": cant * precio,
        "decision": decision, "cant_aprobada": aprob, "importe_aprobado": aprob * precio,
        "frenado": (cant - aprob) * precio, "razon": " ".join(razones), "notas": " ".join(notas),
        "renegociar": renegociar,
    }


def excel(res, fecha, salida, logo):
    wb = Workbook()
    f = lambda **k: Font(name="Open Sans", color=ONYX, **k)
    borde = Border(bottom=Side(style="thin", color=PLATINIUM))
    head_fill = PatternFill("solid", fgColor=ROYAL)

    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False
    if logo and os.path.exists(logo):
        img = Image(logo); img.width, img.height = 160, int(160 * 417 / 1176)
        ws.add_image(img, "B2")
    ws["B6"] = f"Dictamen de órdenes de compra · semana del {fecha}"; ws["B6"].font = f(size=16, bold=True)
    ws["B7"] = "Guardián de compras · para firma de la Dirección General"; ws["B7"].font = f(size=10)
    tot = sum(r["importe"] for r in res); fren = sum(r["frenado"] for r in res)
    ocs = sorted({r["oc"] for r in res})
    filas = [
        ("Órdenes revisadas", len(ocs)), ("Líneas revisadas", len(res)),
        ("Líneas aprobadas", sum(r["decision"] == "APROBAR" for r in res)),
        ("Líneas reducidas", sum(r["decision"] == "REDUCIR" for r in res)),
        ("Líneas rechazadas", sum(r["decision"] == "RECHAZAR" for r in res)),
        ("Monto pedido", tot), ("Monto aprobado", tot - fren), ("Monto que se evitó comprometer", fren),
        ("Ahorro por renegociar precio", sum(r["renegociar"] for r in res)),
    ]
    for k, (a, b) in enumerate(filas, start=9):
        ws.cell(k, 2, a).font = f(size=11)
        c = ws.cell(k, 3, b); c.font = f(size=11, bold=True)
        c.number_format = '"$"#,##0' if isinstance(b, float) else "#,##0"
        if a == "Monto que se evitó comprometer":
            for col in (2, 3): ws.cell(k, col).fill = PatternFill("solid", fgColor=ROYAL_T)
        for col in (2, 3): ws.cell(k, col).border = borde
    ws["B19"] = "Lo que más pesa"; ws["B19"].font = f(size=12, bold=True)
    top = sorted([r for r in res if r["frenado"] > 0], key=lambda r: -r["frenado"])[:8]
    for k, r in enumerate(top, start=20):
        ws.cell(k, 2, f"{r['decision'].title()} · OC {r['oc']} · {r['articulo'][:70]}").font = f(size=10)
        c = ws.cell(k, 3, r["frenado"]); c.font = f(size=10, bold=True); c.number_format = '"$"#,##0'
        ws.cell(k, 4, r["razon"]).font = f(size=9)
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18; ws.column_dimensions["D"].width = 90

    d = wb.create_sheet("Dictamen")
    d.sheet_view.showGridLines = False
    cols = [("OC", "oc", 10), ("Proveedor", "proveedor", 30), ("Artículo", "articulo", 55), ("Tipo", "tipo", 20),
            ("Llegada", "llegada", 12), ("Piezas pedidas", "cant", 12), ("Precio", "precio", 12),
            ("Importe pedido", "importe", 15), ("Dictamen", "decision", 12), ("Piezas aprobadas", "cant_aprobada", 12),
            ("Importe aprobado", "importe_aprobado", 15), ("Evitado", "frenado", 14), ("Razón", "razon", 70), ("Notas", "notas", 60)]
    for j, (t, _, w) in enumerate(cols, start=1):
        c = d.cell(1, j, t); c.font = f(bold=True, size=10); c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
        d.column_dimensions[c.column_letter].width = w
    orden = {"RECHAZAR": 0, "REDUCIR": 1, "APROBAR": 2}
    for i, r in enumerate(sorted(res, key=lambda r: (orden[r["decision"]], -r["frenado"])), start=2):
        fill = {"RECHAZAR": ROYAL_T, "REDUCIR": GINGER_T}.get(r["decision"], SEASHELL if i % 2 else "FFFFFF")
        for j, (_, k, _) in enumerate(cols, start=1):
            c = d.cell(i, j, r[k]); c.font = f(size=9, bold=(k == "decision"))
            c.fill = PatternFill("solid", fgColor=fill); c.border = borde
            c.alignment = Alignment(wrap_text=k in ("razon", "notas", "articulo"), vertical="top")
            if k in ("precio",): c.number_format = '"$"#,##0.00'
            if k in ("importe", "importe_aprobado", "frenado"): c.number_format = '"$"#,##0'
            if k in ("cant", "cant_aprobada"): c.number_format = "#,##0"
    d.freeze_panes = "A2"; d.auto_filter.ref = d.dimensions

    s = wb.create_sheet("Supuestos")
    for k, t in enumerate([
        "Reglas de la Dirección General (15-oct-2026):",
        "1. Básico: cobertura máxima 180 días después de la compra (stock vendible + tránsito + OC). Emergente: sin tope.",
        "2. Bajo pedido, Bajo pedido especial y Bajo Proyecto: solo lo respaldado por pedidos de cliente abiertos.",
        "3. Descontinuado, Hasta Agotar Existencias y Outlet: no se compran.",
        "4. Analizadores: margen ≥10% después de 15% de gastos, con costo puesto = precio OC × 1.12 contra DISTRIBUIDOR IV.",
        "5. Sin mínimos de compra: se marca renegociar si el precio supera el mejor pagado al mismo proveedor en 12 meses.",
        "Inventario vendible: Lindavista, Ciudad de México, Guadalajara, Tijuana, Matriz y Caducidad Corta.",
        "Fuente: NetSuite. Precio de la OC en pesos.",
    ], start=1):
        s.cell(k, 1, t).font = f(size=10, bold=(k == 1))
    s.column_dimensions["A"].width = 120

    os.makedirs(salida, exist_ok=True)
    ruta = os.path.join(salida, f"dictamen-compras-{fecha}.xlsx")
    wb.save(ruta)
    return ruta, len(ocs), tot, fren


if __name__ == "__main__":
    lineas = json.load(open(sys.argv[1]))
    fecha = sys.argv[2]
    salida = sys.argv[3] if len(sys.argv) > 3 else "entregables/dictamenes"
    logo = sys.argv[4] if len(sys.argv) > 4 else os.path.join(os.path.dirname(__file__), "kabla_wordmark_orange.png")
    res = [dictaminar(l) for l in lineas]
    ruta, nocs, tot, fren = excel(res, fecha, salida, logo)
    por_oc = {}
    for r in res:
        o = por_oc.setdefault(r["oc"], {"oc": r["oc"], "proveedor": r["proveedor"], "importe": 0.0, "frenado": 0.0, "decisiones": set()})
        o["importe"] += r["importe"]; o["frenado"] += r["frenado"]; o["decisiones"].add(r["decision"])
    for o in por_oc.values():
        limpia = o["decisiones"] == {"APROBAR"}
        o["usd"] = round(o["importe"] / TIPO_CAMBIO)
        o["firma"] = "Gerente de compras" if limpia and o["usd"] < LIMITE_FIRMA_USD else "Dirección General"
        o["dictamen"] = "OK" if limpia else "REVISAR"
        o["decisiones"] = sorted(o["decisiones"])
    json.dump(sorted(por_oc.values(), key=lambda o: (o["dictamen"] == "OK", -o["frenado"])), open(ruta.replace(".xlsx", "-por-oc.json"), "w"), ensure_ascii=False, indent=1)
    cuenta = {k: sum(r["decision"] == k for r in res) for k in ("APROBAR", "REDUCIR", "RECHAZAR")}
    print(json.dumps({"archivo": ruta, "ordenes": nocs, "lineas": len(res), **cuenta, "monto_pedido": round(tot), "monto_frenado": round(fren),
                      "renegociar": round(sum(r["renegociar"] for r in res)),
                      "firma_gerente": sum(o["firma"] == "Gerente de compras" for o in por_oc.values()),
                      "firma_direccion": sum(o["firma"] == "Dirección General" for o in por_oc.values())}, ensure_ascii=False))
